from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_presentation():
    prs = Presentation()
    
    # Define colors
    bg_color = RGBColor(15, 23, 42) # Deep navy/charcoal
    text_color = RGBColor(241, 245, 249) # Light text
    accent_color = RGBColor(56, 189, 248) # Blue accent
    
    # 1. Title Slide
    slide = prs.slides.add_slide(prs.slide_layouts[6]) # Blank layout
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = bg_color
    
    txBox = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(8), Inches(1.5))
    tf = txBox.text_frame
    p = tf.add_paragraph()
    p.text = "Internship Presentation"
    p.font.size = Pt(44)
    p.font.bold = True
    p.font.color.rgb = text_color
    p.alignment = PP_ALIGN.CENTER
    
    p2 = tf.add_paragraph()
    p2.text = "[Your Name]\n[Internship Role/Domain]\n[Company Name]\n[Start Month – End Month, Year]"
    p2.font.size = Pt(24)
    p2.font.color.rgb = accent_color
    p2.alignment = PP_ALIGN.CENTER
    
    p3 = tf.add_paragraph()
    p3.text = "Full-Stack E-Commerce Platform & Real-Time Collaborative Workspace"
    p3.font.size = Pt(18)
    p3.font.color.rgb = text_color
    p3.alignment = PP_ALIGN.CENTER
    
    # Helper for content slides
    def add_content_slide(title, bullets):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        bg = slide.background
        fill = bg.fill
        fill.solid()
        fill.fore_color.rgb = bg_color
        
        # Title
        title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(9), Inches(1))
        tf_title = title_box.text_frame
        p_title = tf_title.add_paragraph()
        p_title.text = title
        p_title.font.size = Pt(36)
        p_title.font.bold = True
        p_title.font.color.rgb = text_color
        
        # Divider line
        line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.5), Inches(1.3), Inches(9), Inches(0.05))
        line.fill.solid()
        line.fill.fore_color.rgb = accent_color
        line.line.color.rgb = accent_color
        
        # Bullets
        content_box = slide.shapes.add_textbox(Inches(0.5), Inches(1.6), Inches(9), Inches(5))
        tf_content = content_box.text_frame
        for bullet in bullets:
            p = tf_content.add_paragraph()
            p.text = "• " + bullet
            p.font.size = Pt(24)
            p.font.color.rgb = text_color
            p.space_after = Pt(14)
            
        # Footer
        footer_box = slide.shapes.add_textbox(Inches(0.5), Inches(7), Inches(9), Inches(0.5))
        tf_footer = footer_box.text_frame
        p_footer = tf_footer.add_paragraph()
        slide_num = len(prs.slides)
        p_footer.text = f"[Your Name] | [Internship Title] | Slide {slide_num}"
        p_footer.font.size = Pt(12)
        p_footer.font.color.rgb = RGBColor(148, 163, 184)
        
        return slide

    # 2. About Company
    add_content_slide("About Company", [
        "[FILL IN: Company Name]",
        "[FILL IN: Industry/Domain]",
        "[FILL IN: Company Size]",
        "[FILL IN: What the company does]",
        "[FILL IN: Team/Department I worked in]"
    ])

    # 3. Internship Objective
    add_content_slide("Internship Objective", [
        "Learn and apply full-stack web development",
        "Build complete end-to-end applications",
        "Implement secure database design",
        "Deliver E-Commerce & Project Management apps"
    ])

    # 4. Technology Stack
    slide4 = prs.slides.add_slide(prs.slide_layouts[6])
    bg4 = slide4.background
    bg4.fill.solid()
    bg4.fill.fore_color.rgb = bg_color
    
    tbox4 = slide4.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(9), Inches(1))
    tbox4.text_frame.text = "Technology Stack"
    tbox4.text_frame.paragraphs[0].font.size = Pt(36)
    tbox4.text_frame.paragraphs[0].font.bold = True
    tbox4.text_frame.paragraphs[0].font.color.rgb = text_color
    line4 = slide4.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.5), Inches(1.3), Inches(9), Inches(0.05))
    line4.fill.solid()
    line4.fill.fore_color.rgb = accent_color
    line4.line.color.rgb = accent_color
    
    cols = 2
    rows = 6
    table = slide4.shapes.add_table(rows, cols, Inches(0.5), Inches(1.8), Inches(9), Inches(4)).table
    stack_data = [
        ("Category", "Technologies"),
        ("Frontend", "HTML5, CSS3, JavaScript ES6, React 18"),
        ("Backend/Logic", "Node.js, Express.js"),
        ("Database", "MongoDB Atlas, JSON Datastore"),
        ("Real-time", "Socket.IO, WebSockets"),
        ("Tools/Platforms", "Vite, Git, GitHub, Docker")
    ]
    for r in range(rows):
        for c in range(cols):
            cell = table.cell(r, c)
            cell.text = stack_data[r][c]
            for p in cell.text_frame.paragraphs:
                p.font.size = Pt(20)
                p.font.color.rgb = text_color
                if r == 0:
                    p.font.bold = True
                    p.font.color.rgb = accent_color
                    
    footer_box4 = slide4.shapes.add_textbox(Inches(0.5), Inches(7), Inches(9), Inches(0.5))
    footer_box4.text_frame.text = f"[Your Name] | [Internship Title] | Slide 4"
    footer_box4.text_frame.paragraphs[0].font.size = Pt(12)
    footer_box4.text_frame.paragraphs[0].font.color.rgb = RGBColor(148, 163, 184)

    # 5. Key Learnings
    add_content_slide("Key Learnings", [
        "Structuring scalable full-stack applications",
        "Building responsive glassmorphism UIs",
        "Implementing secure JWT authentication",
        "Real-time WebSocket synchronization",
        "Using Git effectively for collaboration"
    ])

    # 6. Challenges Faced
    add_content_slide("Challenges Faced", [
        "Resolving environment setup issues",
        "Structuring robust project architecture",
        "Validating cart prices securely server-side",
        "Managing real-time Socket.IO room states",
        "Cross-browser UI consistency for glassmorphism"
    ])

    # 7. Results and Conclusions
    add_content_slide("Results and Conclusions", [
        "Delivered AroraCart E-Commerce Platform",
        "Delivered TaskPulse Real-Time Workspace",
        "Implemented robust authentication & cart logic",
        "Significantly improved full-stack dev skills"
    ])

    # 8. Completion Certificate
    slide8 = prs.slides.add_slide(prs.slide_layouts[6])
    bg8 = slide8.background
    bg8.fill.solid()
    bg8.fill.fore_color.rgb = bg_color
    
    tbox8 = slide8.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(9), Inches(1))
    tbox8.text_frame.text = "Completion Certificate"
    tbox8.text_frame.paragraphs[0].font.size = Pt(36)
    tbox8.text_frame.paragraphs[0].font.bold = True
    tbox8.text_frame.paragraphs[0].font.color.rgb = text_color
    line8 = slide8.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.5), Inches(1.3), Inches(9), Inches(0.05))
    line8.fill.solid()
    line8.fill.fore_color.rgb = accent_color
    line8.line.color.rgb = accent_color
    
    placeholder = slide8.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(2), Inches(2), Inches(6), Inches(4))
    placeholder.fill.solid()
    placeholder.fill.fore_color.rgb = RGBColor(51, 65, 85)
    placeholder.line.color.rgb = text_color
    ptf = placeholder.text_frame
    ptf.text = "[INSERT CERTIFICATE IMAGE HERE]"
    ptf.paragraphs[0].font.size = Pt(24)
    ptf.paragraphs[0].font.color.rgb = text_color
    ptf.paragraphs[0].alignment = PP_ALIGN.CENTER
    
    cap = slide8.shapes.add_textbox(Inches(0.5), Inches(6.2), Inches(9), Inches(0.5))
    cap.text_frame.text = "Certificate of Completion — [Company Name]"
    cap.text_frame.paragraphs[0].font.size = Pt(18)
    cap.text_frame.paragraphs[0].font.color.rgb = text_color
    cap.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    
    footer_box8 = slide8.shapes.add_textbox(Inches(0.5), Inches(7), Inches(9), Inches(0.5))
    footer_box8.text_frame.text = f"[Your Name] | [Internship Title] | Slide 8"
    footer_box8.text_frame.paragraphs[0].font.size = Pt(12)
    footer_box8.text_frame.paragraphs[0].font.color.rgb = RGBColor(148, 163, 184)

    # 9. Thank You Slide
    slide9 = prs.slides.add_slide(prs.slide_layouts[6])
    bg9 = slide9.background
    bg9.fill.solid()
    bg9.fill.fore_color.rgb = bg_color
    
    txBox9 = slide9.shapes.add_textbox(Inches(1), Inches(3), Inches(8), Inches(2))
    tf9 = txBox9.text_frame
    p9 = tf9.add_paragraph()
    p9.text = "Thank You"
    p9.font.size = Pt(60)
    p9.font.bold = True
    p9.font.color.rgb = text_color
    p9.alignment = PP_ALIGN.CENTER
    
    p9_2 = tf9.add_paragraph()
    p9_2.text = "[Your Name] | [Contact Info: Email/LinkedIn/GitHub]"
    p9_2.font.size = Pt(24)
    p9_2.font.color.rgb = accent_color
    p9_2.alignment = PP_ALIGN.CENTER
    
    prs.save('d:/codealpha/codealpha_task1-main/Internship_Presentation.pptx')

if __name__ == '__main__':
    create_presentation()
